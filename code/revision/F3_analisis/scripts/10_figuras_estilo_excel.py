"""Figuras del paper con el diseño de los gráficos de Systematic_Review_MQ_Analysis.xlsx.

Replica el estilo del Excel original: tema Office 2007–2010 (accent1–6 y sus variantes de luminosidad),
columnas con degradado vertical y un color por barra, etiquetas blancas en negrita dentro de la barra,
fondo gris con degradado radial y borde redondeado, tortas con porcentajes en recuadro oscuro y línea suavizada
con marcadores grandes para MQ5. Los datos se leen del Excel v2, que es la fuente única.

MQ6 se dibuja como barras: el mapa (chartEx de Excel) no se puede reproducir sin datos geográficos locales;
el mapa actualizado está en el Excel y se puede exportar desde ahí.

Entrada: Systematic_Review_MQ_Analysis_v2.xlsx
Salida:  revision/F3_analisis/figuras/estilo_excel/*.png (300 dpi) y *.pdf
Uso:     python revision/F3_analisis/scripts/10_figuras_estilo_excel.py
"""
import shutil
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from matplotlib.patches import FancyBboxPatch
from scipy.interpolate import PchipInterpolator

ROOT = Path(__file__).resolve().parents[3]
XLSX = ROOT / "Systematic_Review_MQ_Analysis_v2.xlsx"
OUT = ROOT / "revision" / "F3_analisis" / "figuras" / "estilo_excel"

ACCENTS = ["#4F81BD", "#C0504D", "#9BBB59", "#8064A2", "#4BACC6", "#F79646"]
TEXT = "#595959"
plt.rcParams.update({"font.family": "Calibri", "font.size": 11})


def rgb(h):
    return np.array([int(h[i:i + 2], 16) for i in (1, 3, 5)]) / 255


def variant(h, k):
    """Ciclo de colores del Excel: base, lumMod 60 %, lumMod 80 % + lumOff 20 %, lumMod 80 %."""
    c = rgb(h)
    return [c, c * 0.6, c * 0.8 + 0.2, c * 0.8][k % 4]


def color(i):
    return variant(ACCENTS[i % 6], i // 6)


def background(fig):
    """Degradado radial gris (blanco al centro) con borde redondeado, como el área del gráfico de Excel."""
    ax = fig.add_axes([0, 0, 1, 1], zorder=-10)
    y, x = np.mgrid[-1:1:400j, -1:1:600j]
    r = np.clip(np.sqrt(x ** 2 + (y * 0.9) ** 2) / 1.35, 0, 1)
    img = np.dstack([1 - 0.22 * r] * 3)
    ax.imshow(img, extent=[0, 1, 0, 1], aspect="auto")
    ax.add_patch(FancyBboxPatch((0.004, 0.006), 0.992, 0.988, boxstyle="round,pad=0,rounding_size=0.015",
                                fc="none", ec="#bfbfbf", lw=1.2, transform=ax.transAxes))
    ax.set_axis_off()


def gradient_bar(ax, x, h, c, width=0.68):
    """Columna con degradado vertical (100 % → 84 % de luminosidad) y sombra suave."""
    if h <= 0:
        return
    ax.bar(x + 0.06, h * 0.985, width, color="black", alpha=0.10, zorder=1)
    grad = np.linspace(1, 0.84, 64).reshape(-1, 1, 1) * np.asarray(c).reshape(1, 1, 3)
    ax.imshow(grad, extent=[x - width / 2, x + width / 2, 0, h], aspect="auto", zorder=2, origin="upper")


def save(fig, name):
    OUT.mkdir(parents=True, exist_ok=True)
    fig.savefig(OUT / f"{name}.png", dpi=300)
    fig.savefig(OUT / f"{name}.pdf")
    # JPG para el manuscrito: pdflatex lo inserta sin recodificar (los PNG con degradados alargan la compilación)
    fig.savefig(OUT / f"{name}.jpg", dpi=200, pil_kwargs={"quality": 90, "optimize": True})
    plt.close(fig)
    manuscript_images = ROOT / "Review_Systematic" / "images"  # el .tex usa estos JPG: se copian siempre
    if manuscript_images.exists():
        shutil.copy(OUT / f"{name}.jpg", manuscript_images / f"{name}.jpg")


def read_block(ws, start=4, col_label=1, col_val=2):
    labels, vals = [], []
    r = start
    while True:
        lab, v = ws.cell(r, col_label).value, ws.cell(r, col_val).value
        if lab is None or not isinstance(v, (int, float)) or str(lab) == "Total":
            break
        labels.append(str(lab))
        vals.append(v)
        r += 1
    return labels, vals


def column_chart(labels, vals, title, name, ylabel="Number of articles", wrap=48):
    n = len(labels)
    fig = plt.figure(figsize=(max(7.5, 0.55 * n + 3), 5.6))
    background(fig)
    longest = max(len(l) for l in labels)
    left = min(0.20, 0.06 + 0.0042 * min(len(labels[0]), wrap))  # deja sitio a la primera etiqueta rotada
    ax = fig.add_axes([left, 0.36 if longest > 18 else 0.22, 0.97 - left, 0.52])
    ax.set_facecolor("none")
    for i, v in enumerate(vals):
        gradient_bar(ax, i, v, color(i))
        inside = v >= max(vals) * 0.09  # barra baja: la etiqueta va encima, como hace Excel cuando no cabe
        ax.text(i, v - max(vals) * 0.035 if inside else v + max(vals) * 0.015, f"{int(v)}", ha="center",
                va="top" if inside else "bottom", color="white" if inside else "#404040",
                fontsize=11, fontweight="bold", zorder=3)
    ax.set_xlim(-0.6, n - 0.4)
    ax.set_ylim(0, max(vals) * 1.05)
    short = [l if len(l) <= wrap else l[:wrap - 1] + "…" for l in labels]
    ax.set_xticks(range(n), short, rotation=45, ha="right", rotation_mode="anchor", color=TEXT, fontsize=10)
    ax.tick_params(axis="x", length=0)
    ax.set_yticks([])
    for s in ax.spines.values():
        s.set_visible(False)
    ax.set_ylabel(ylabel, color="#404040", fontweight="bold", fontsize=11)
    fig.suptitle(title, color=TEXT, fontsize=20, y=0.965)
    save(fig, name)


def pie_chart(labels, vals, title, name):
    fig = plt.figure(figsize=(8.0, 5.8))
    background(fig)
    ax = fig.add_axes([0.03, 0.04, 0.62, 0.92])
    cols = [color(i) for i in range(len(vals))]
    wedges, _ = ax.pie(vals, colors=cols, startangle=90, counterclock=False,
                       wedgeprops=dict(edgecolor="none"))
    total = sum(vals)
    for w, v in zip(wedges, vals):
        if v / total < 0.015:
            continue
        ang = np.deg2rad((w.theta1 + w.theta2) / 2)
        ax.text(0.62 * np.cos(ang), 0.62 * np.sin(ang), f"{100 * v / total:.0f}%", ha="center", va="center",
                color="white", fontsize=12, fontweight="bold",
                bbox=dict(boxstyle="square,pad=0.25", fc="#404040", ec="none", alpha=0.92))
    ax.set_aspect("equal")
    fig.legend(wedges, [f"{l} ({v})" for l, v in zip(labels, vals)], loc="center left",
               bbox_to_anchor=(0.66, 0.45), frameon=True, facecolor="#f2f2f2", edgecolor="none", fontsize=11,
               labelcolor="#404040")
    fig.text(0.97, 0.88, title, ha="right", va="center", fontsize=20, fontweight="bold", color="#404040")
    save(fig, name)


def line_chart(years, vals, title, name, partial=None):
    fig = plt.figure(figsize=(11, 5.4))
    background(fig)
    ax = fig.add_axes([0.06, 0.12, 0.92, 0.72])
    ax.set_facecolor("none")
    x = np.arange(len(years))
    xs = np.linspace(0, len(years) - 1, 300)
    ys = PchipInterpolator(x, vals)(xs)  # monótona: no inventa valles ni picos entre puntos
    dark, light = rgb("#8B2C2C"), rgb("#F2D0D0")
    pal = [dark + (light - dark) * t for t in np.linspace(0, 1, len(years))]
    for i in range(len(xs) - 1):
        t = xs[i] / (len(years) - 1)
        ax.plot(xs[i:i + 2], ys[i:i + 2], color=dark + (light - dark) * t, lw=3, solid_capstyle="round", zorder=2)
    for i, (xi, v) in enumerate(zip(x, vals)):
        ax.scatter(xi, v, s=620, color=pal[i], zorder=3, edgecolor="none")
        ax.text(xi, v, f"{int(v)}", ha="center", va="center", color="white", fontsize=11, fontweight="bold", zorder=4)
    labels = [f"{y}*" if y == partial else str(y) for y in years]
    ax.set_xticks(x, labels, color=TEXT, fontsize=11)
    ax.tick_params(axis="both", length=0)
    ax.set_yticks(np.arange(0, max(vals) * 1.25 + 1, max(1, round(max(vals) / 5))))
    ax.set_yticklabels([])
    ax.yaxis.grid(True, color="#a6a6a6", lw=0.8)
    ax.set_ylim(0, max(vals) * 1.22)
    ax.set_xlim(-0.5, len(years) - 0.5)
    for k in ("top", "right", "left"):
        ax.spines[k].set_visible(False)
    ax.spines["bottom"].set_color("#262626")
    ax.spines["bottom"].set_linewidth(2)
    ax.set_ylabel("Number of articles", color="#404040", fontweight="bold")
    fig.suptitle(title, color="#404040", fontsize=22, fontweight="bold", y=0.96)
    if partial:
        fig.text(0.98, 0.02, f"*{partial}: partial year", ha="right", color=TEXT, fontsize=9)
    save(fig, name)


def main():
    try:
        wb = openpyxl.load_workbook(XLSX)
    except PermissionError:  # abierto en Excel: se lee una copia de la versión guardada (Copy-Item permite leerlo)
        import subprocess
        import tempfile
        tmp = Path(tempfile.gettempdir()) / "mq_v2_copia.xlsx"
        subprocess.run(["powershell", "-NoProfile", "-Command",
                        f"Copy-Item -LiteralPath '{XLSX}' -Destination '{tmp}' -Force -ErrorAction Stop"], check=True)
        wb = openpyxl.load_workbook(tmp)
        print("Nota: el Excel MQ v2 está abierto; se leyó la versión guardada en disco.")
    lab, val = read_block(wb["MQ1_Database"])
    pie_chart(lab, val, "Database Frequency", "mq01_database")
    lab, val = read_block(wb["MQ2_Reviewed_vs_Selected"])
    column_chart(lab, val, "Study Selection Flow", "mq02_selection_flow")
    lab, val = read_block(wb["MQ2_Reviewed_vs_Selected"], start=12)
    column_chart(lab, val, "Included Studies by Predictive Task", "mq02_studies_by_task")
    for sheet, title, name in [("MQ3_Authors", "Most Common Authors", "mq03_authors"),
                               ("MQ4_Journals", "Most Used Journals and Venues", "mq04_journals"),
                               ("MQ6_Countries", "Most Frequent Countries", "mq06_countries"),
                               ("MQ7_Datasets", "Most Used Dataset Families", "mq07_datasets"),
                               ("MQ8_RegressionTypes", "Models Used for Continuous Outcomes (Density / Count)", "mq08_continuous_models"),
                               ("MQ9_EvalMetrics", "Evaluation Metric Groups", "mq09_metrics"),
                               ("MQ11_MLModels", "Most Used Machine Learning Models", "mq11_ml_models"),
                               ("MQ12_BioInspired", "Most Used Bio-inspired Algorithms", "mq12_bio_algorithms")]:
        lab, val = read_block(wb[sheet])
        column_chart(lab, val, title, name)
    lab, val = read_block(wb["MQ5_Year"])
    line_chart([int(float(l)) for l in lab], val, "Articles per Year", "mq05_years", partial=2026)
    lab, val = read_block(wb["MQ10_HyperparamTuning"])
    pie_chart(lab, val, "Hyperparameter Tuning Technique", "mq10_hyperparameter")
    print("Figuras:", sorted(p.name for p in OUT.glob("*.png")))


if __name__ == "__main__":
    main()
