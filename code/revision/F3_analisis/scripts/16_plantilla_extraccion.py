"""Genera la plantilla de extracción a texto completo para los estudios añadidos por snowballing / control.

Hojas:
  Instrucciones          cómo completarla, leyenda de códigos y una fila de ejemplo (estudio P211 de la búsqueda)
  Extraccion_densidad    los estudios de densidad añadidos (corpus C y A): OBLIGATORIA
  Extraccion_recuento    los estudios de recuento (sólo síntesis B): OPCIONAL
Las celdas amarillas son las que completa el revisor; las grises son de referencia (lo codificado hoy desde
título y resumen) y no se leen.

No sobrescribe una plantilla existente (para no perder extracciones): usar --forzar para regenerarla.
Uso: python revision/F3_analisis/scripts/16_plantilla_extraccion.py [--forzar]
"""
import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
import fulltext as ft  # noqa: E402

REV = Path(__file__).resolve().parents[2]
HEAD = PatternFill("solid", fgColor="1F4E78")
INPUT = PatternFill("solid", fgColor="FFF2CC")
REF = PatternFill("solid", fgColor="EDEDED")
FIXED = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
F = "Arial"

FIXED_COLS = ["ID", "Título", "Autores", "Año", "Venue", "DOI", "Enlace", "Tarea actual"]
INPUT_COLS = [h for h in ft.COLUMNS if h not in FIXED_COLS]
REF_COLS = ["Ref: datasets (resumen)", "Ref: modelos (resumen)", "Ref: métricas (resumen)",
            "Ref: HPO (resumen)", "Ref: bio-inspirados (resumen)"]
WIDTH = {"ID": 7, "Título": 48, "Autores": 28, "Año": 6, "Venue": 26, "DOI": 26, "Enlace": 30, "Tarea actual": 11,
         "Revisor": 14, "Fecha": 11, "Tarea confirmada": 12, "Medida de tamaño (si es densidad)": 22,
         "Datasets": 34, "Modelos de ML / estadísticos": 34, "Algoritmos bio-inspirados y su rol": 30,
         "Hiperparámetros (técnica de ajuste)": 30, "Métricas de evaluación": 30,
         "Diseño experimental / metodología": 34, "Conclusiones": 34, "Decisión final": 12,
         "Motivo de exclusión": 14, "Notas": 30}
YESNO = [h for h in INPUT_COLS if h.startswith("¿")]
QCOLS = [h for h in INPUT_COLS if h.startswith("Q")]


def header(ws, cols):
    for j, h in enumerate(cols, 1):
        c = ws.cell(1, j, h)
        c.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BOX
        ws.column_dimensions[get_column_letter(j)].width = WIDTH.get(h, 13 if h in YESNO or h in QCOLS else 24)
    ws.row_dimensions[1].height = 48
    ws.freeze_panes = "C2"


def validations(ws, cols, n):
    last = n + 1

    def add(values, names):
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Valor no válido", error="Elija un valor de la lista.")
        ws.add_data_validation(dv)
        for h in names:
            col = get_column_letter(cols.index(h) + 1)
            dv.add(f"{col}2:{col}{last}")

    add(["Sí", "No"], YESNO)
    add(["+1", "0", "-1"], QCOLS)
    add(ft.TASKS, ["Tarea confirmada"])
    add(["Incluir", "Excluir"], ["Decisión final"])
    add(ft.EXCLUSION_CODES, ["Motivo de exclusión"])


def sheet(wb, name, studies, ref):
    ws = wb.create_sheet(name)
    cols = FIXED_COLS + INPUT_COLS + REF_COLS
    header(ws, cols)
    for i, r in enumerate(studies.itertuples(), 2):
        doi = r.doi
        rf = ref.get(r.id, {})
        values = {"ID": r.id, "Título": r.title, "Autores": r.autores, "Año": r.year, "Venue": r.venue, "DOI": doi,
                  "Enlace": f"https://doi.org/{doi}" if doi else "", "Tarea actual": r.tarea,
                  REF_COLS[0]: rf.get("datasets", ""), REF_COLS[1]: rf.get("modelos", ""),
                  REF_COLS[2]: rf.get("metricas", ""), REF_COLS[3]: rf.get("hpo_tecnica", ""),
                  REF_COLS[4]: rf.get("bio", "")}
        for j, h in enumerate(cols, 1):
            c = ws.cell(i, j, values.get(h, ""))
            c.font = Font(name=F, size=9, italic=h in REF_COLS, color="595959" if h in REF_COLS else "000000")
            c.fill = FIXED if h in FIXED_COLS else REF if h in REF_COLS else INPUT
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BOX
            if h == "Enlace" and values["Enlace"]:
                c.hyperlink = values["Enlace"]
                c.font = Font(name=F, size=9, color="0563C1", underline="single")
    validations(ws, cols, len(studies))
    return ws


def instructions(wb, example):
    ws = wb.active
    ws.title = "Instrucciones"
    lines = [
        ("Extracción a texto completo: estudios añadidos por snowballing y conjunto de control", "title"),
        ("", ""),
        ("Qué hacer", "h"),
        ("1. Abra el artículo con el enlace DOI y léalo completo.", ""),
        ("2. Complete SOLO las celdas amarillas, con el mismo criterio de la ficha original de la búsqueda.", ""),
        ("3. Las columnas grises (Ref: …) muestran lo que se codificó hoy desde el título y el resumen: úselas solo como orientación; no se leen.", ""),
        ("4. Al terminar, ponga la 'Decisión final' (Incluir / Excluir). Solo las filas con decisión se usan en el análisis.", ""),
        ("5. Si excluye un estudio, indique el 'Motivo de exclusión' (E1–E7).", ""),
        ("6. Luego ejecute el pipeline (ver README_F3.md): 02 → 06 → 07 → 09 → 10 → 11 → 14 → 15.", ""),
        ("", ""),
        ("Leyenda", "h"),
        ("Hoja Extraccion_densidad: OBLIGATORIA (estudios del mapeo C y de la síntesis de densidad A).", ""),
        ("Hoja Extraccion_recuento: opcional (estudios de recuento, solo para la síntesis B).", ""),
        ("Tarea confirmada: DENS = densidad de defectos · COUNT = número de defectos · CLASS = clasificación de propensión · "
         "EFFORT = ranking effort-aware · EMP = modelo estadístico empírico · REL = fiabilidad · LOCALIZ = localización de fallos. "
         "Déjela vacía si la tarea actual es correcta.", ""),
        ("Motivo de exclusión: E1 no accesible · E2 no trata defectos de software · E3 sin datasets · E4 sin modelo predictivo · "
         "E5 sin métricas de evaluación · E6 fuera de ingeniería de software · E7 estudio secundario.", ""),
        ("Calidad (Q1–Q4): +1 = Sí · 0 = Parcialmente · -1 = No.", ""),
        ("Medida de tamaño: denominador de la densidad (LOC, KLOC, puntos de función…), o 'no reportada'.", ""),
        ("Hiperparámetros: nombre la TÉCNICA de ajuste (grid search, random search, bayesiana, GA/PSO…). Si solo se fijan valores sin "
         "búsqueda, escriba 'valores fijos, sin búsqueda'. Los coeficientes o pesos del modelo no son hiperparámetros.", ""),
        ("Bio-inspirados: algoritmo y su rol (selección de características / ajuste de hiperparámetros / entrenamiento del modelo).", ""),
        ("", ""),
        ("Ejemplo (estudio P211 de la búsqueda, ya extraído en la ficha original)", "h"),
    ]
    r = 1
    for text, kind in lines:
        c = ws.cell(r, 1, text)
        c.font = Font(name=F, size=14 if kind == "title" else 11, bold=kind in ("title", "h"),
                      color="1F4E78" if kind in ("title", "h") else "000000")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110
    for k, v in example:
        a, b = ws.cell(r, 1, k), ws.cell(r, 2, v)
        a.font, b.font = Font(name=F, size=10, bold=True), Font(name=F, size=10)
        a.fill, b.fill = FIXED, INPUT
        b.alignment = Alignment(wrap_text=True, vertical="top")
        a.border = b.border = BOX
        r += 1
    for rr in range(1, r):
        if ws.cell(rr, 2).value is None:
            ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=2)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--forzar", action="store_true")
    a = ap.parse_args()
    if ft.PATH.exists() and not a.forzar:
        sys.exit(f"Ya existe {ft.PATH.name}; no se sobrescribe (podría tener extracciones). Use --forzar para regenerarla.")
    cod = pd.read_csv(REV / "F3_analisis" / "corpus_codificado.csv", keep_default_na=False, dtype=str)
    m = pd.read_excel(REV / "F1_corpus" / "extraction_matrix_v2.xlsx", sheet_name="Corpus_v2",
                      keep_default_na=False, dtype=str)
    x = cod.merge(m[["id", "title", "doi"]], on="id", how="left")
    x = x[x["fuente"] != "busqueda_bd"].sort_values("id")
    ref = x.set_index("id")[["datasets", "modelos", "metricas", "hpo_tecnica", "bio"]].to_dict("index")
    p211 = m.set_index("id").loc["P211"]
    example = [("ID", "P211"), ("Título", p211.title), ("Tarea confirmada", "(vacía: DENS es correcta)"),
               ("Medida de tamaño (si es densidad)", "LOC (defects per line of code)"),
               ("Datasets", p211.datasets_text), ("Modelos de ML / estadísticos", p211.ml_text),
               ("Algoritmos bio-inspirados y su rol", "No"), ("Hiperparámetros (técnica de ajuste)", "Grid search"),
               ("Métricas de evaluación", p211.metrics_text), ("Diseño experimental / metodología", p211.design_text[:300]),
               ("Q1–Q4", "+1 · 0 · +1 · +1"), ("Decisión final", "Incluir")]
    wb = Workbook()
    instructions(wb, example)
    sheet(wb, "Extraccion_densidad", x[x.en_C == "True"], ref)
    sheet(wb, "Extraccion_recuento", x[x.en_C != "True"], ref)
    wb.save(ft.PATH)
    print(f"Plantilla: {ft.PATH} · densidad {int((x.en_C == 'True').sum())} · recuento {int((x.en_C != 'True').sum())}")


if __name__ == "__main__":
    main()
